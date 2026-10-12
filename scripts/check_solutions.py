"""Verify best solutions with the course's spreadsheet solution checker.

Usage:  conda run -n ec python scripts/check_solutions.py lab1

For every method, the best TSPA and TSPB tours are pasted into a copy of
problem/Solution checker.xlsx ("List of nodes" column, F3:F102), the workbook
is recalculated by LibreOffice in headless mode, and the objective is read back
from cell K2. The result is compared with the value reported in summary.csv.

Writes labN/results/checker.csv.

Requires LibreOffice (soffice). The checker itself is the lecturer's file and
is not in the repository; put it in problem/ before running.
"""

import csv
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CHECKER = ROOT / "problem" / "Solution checker.xlsx"
SOFFICE_CANDIDATES = ["soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
FIRST_ROW = 3  # node list occupies F3:F102; F103 = F3 closes the cycle


def find_soffice() -> str:
    for candidate in SOFFICE_CANDIDATES:
        if shutil.which(candidate) or Path(candidate).exists():
            return candidate
    sys.exit("LibreOffice (soffice) not found")


def recalculate(soffice: str, src: Path, out_dir: Path) -> Path:
    """Round-trips the workbook through LibreOffice so all formulas get values."""
    # A private profile keeps this working even when LibreOffice is open in the GUI.
    profile = (out_dir / "lo-profile").as_uri()
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless", "--calc",
         "--convert-to", "xlsx", "--outdir", str(out_dir / "out"), str(src)],
        check=True, capture_output=True,
    )
    return out_dir / "out" / src.name


def main(lab: str):
    if not CHECKER.exists():
        sys.exit(f"checker not found: {CHECKER}")
    soffice = find_soffice()
    results = ROOT / lab / "results"

    reported = {}
    with open(results / "summary.csv") as f:
        for row in csv.DictReader(f):
            reported[(row["instance"], row["method"])] = int(row["min"])

    tours: dict[str, dict[str, list[int]]] = {}
    for path in sorted((results / "best").glob("*.txt")):
        instance, method = path.stem.split("_", 1)
        tours.setdefault(method, {})[instance] = [int(v) for v in path.read_text().split()]

    rows = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for method, by_instance in tours.items():
            wb = openpyxl.load_workbook(CHECKER)
            for instance, tour in by_instance.items():
                ws = wb[instance]
                if len(tour) != 100:
                    sys.exit(f"{instance} {method}: checker expects 100 nodes, got {len(tour)}")
                for i, node in enumerate(tour):
                    ws.cell(row=FIRST_ROW + i, column=6, value=node)  # column F
            src = tmp / f"{method}.xlsx"
            wb.save(src)

            values = openpyxl.load_workbook(recalculate(soffice, src, tmp), data_only=True)
            for instance in by_instance:
                checker_value = values[instance]["K2"].value
                expected = reported[(instance, method)]
                ok = checker_value == expected
                rows.append((instance, method, expected, checker_value, ok))
                print(f"{instance} {method:<13} reported {expected:>7}  checker {checker_value!s:>7}  "
                      f"{'OK' if ok else 'MISMATCH'}")

    rows.sort()
    with open(results / "checker.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["instance", "method", "reported", "checker", "match"])
        w.writerows(rows)

    if not all(r[4] for r in rows):
        sys.exit("some solutions do not match the checker")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
